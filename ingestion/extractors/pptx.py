"""PPTX multimodal extraction (text + images)."""

from __future__ import annotations

from pathlib import Path
from typing import List

from ingestion.chunker import split_by_size
from ingestion.chunk_types import RawChunk
from multimodal.base import Modality


def _load_pptx():
    try:
        from pptx import Presentation  # type: ignore
        from pptx.enum.shapes import MSO_SHAPE_TYPE  # type: ignore
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise RuntimeError("python-pptx is required for PPTX extraction") from exc
    return Presentation, MSO_SHAPE_TYPE


def extract_pptx_chunks(path: Path, *, chunk_size: int = 800) -> List[RawChunk]:
    Presentation, MSO_SHAPE_TYPE = _load_pptx()
    pres = Presentation(path)
    chunks: List[RawChunk] = []

    for slide_idx, slide in enumerate(pres.slides, start=1):
        # Aggregate text from all text-containing shapes on the slide
        texts = []
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            text = shape.text.strip()
            if text:
                texts.append(text)
        if texts:
            slide_text = "\n".join(texts)
            for idx, chunk_text in enumerate(split_by_size(slide_text, max_chars=chunk_size)):
                chunk_id = f"{path.stem}-s{slide_idx}-t{idx}"
                chunks.append(
                    RawChunk(
                        chunk_id=chunk_id,
                        modality=Modality.TEXT,
                        content=chunk_text,
                        metadata={
                            "modality": Modality.TEXT.value,
                            "slide": slide_idx,
                            "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        },
                    )
                )

        # Collect images
        for shape_idx, shape in enumerate(slide.shapes, start=1):
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            image = shape.image
            blob = image.blob
            if not blob:
                continue
            chunk_id = f"{path.stem}-s{slide_idx}-img{shape_idx}"
            chunks.append(
                RawChunk(
                    chunk_id=chunk_id,
                    modality=Modality.IMAGE,
                    content=blob,
                    metadata={
                        "modality": Modality.IMAGE.value,
                        "slide": slide_idx,
                        "image_ext": image.ext,
                        "mime": image.mime_type,
                    },
                )
            )

    return chunks


__all__ = ["extract_pptx_chunks"]
